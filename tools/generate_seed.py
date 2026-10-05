import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def sql_text(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def iso(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    raise ValueError(f"Fecha no válida: {value!r}")


def n(value):
    return format(float(value), ".15g")


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    wb = load_workbook(source, data_only=True, read_only=False)
    heating = wb["Lecturas Cal"]
    cooling = wb["Lecturas Frío"]
    water = wb["Lecturas Agua"]
    invoices = wb["Facturas reales"]
    summary = wb["Resumen periodo"]
    prices = wb["Cálculos precios"]

    lines = ["PRAGMA foreign_keys = ON;", ""]
    dwelling_ids = {}
    for row in range(2, 81):
        address = heating.cell(row, 1).value
        short_name = str(address).split(" - ")[-1]
        dwelling_id = row - 1
        dwelling_ids[address] = dwelling_id
        lines.append(
            "INSERT OR IGNORE INTO dwellings (id, address, short_name, sort_order) VALUES "
            f"({dwelling_id}, {sql_text(address)}, {sql_text(short_name)}, {dwelling_id});"
        )

    lines.append("")
    date_ids = {}
    for column in range(2, heating.max_column + 1):
        value = heating.cell(1, column).value
        if not isinstance(value, (date, datetime)):
            continue
        date_id = len(date_ids) + 1
        date_value = iso(value)
        date_ids[date_value] = date_id
        lines.append(
            "INSERT OR IGNORE INTO reading_dates (id, reading_date) VALUES "
            f"({date_id}, {sql_text(date_value)});"
        )

    lines.append("")
    service_sheets = [("heating", heating), ("cooling", cooling), ("water", water)]
    for service, sheet in service_sheets:
        for column in range(2, heating.max_column + 1):
            header = heating.cell(1, column).value
            if not isinstance(header, (date, datetime)):
                continue
            date_id = date_ids[iso(header)]
            for row in range(2, 81):
                value = sheet.cell(row, column).value
                if value is None:
                    continue
                lines.append(
                    "INSERT OR REPLACE INTO readings (dwelling_id, reading_date_id, service, value) VALUES "
                    f"({row - 1}, {date_id}, {sql_text(service)}, {n(value)});"
                )

    lines.append("")
    invoice_types = {"Luz": "electricity", "Agua": "water"}
    for row in range(2, invoices.max_row + 1):
        raw_type = str(invoices.cell(row, 1).value).strip()
        lines.append(
            "INSERT OR IGNORE INTO invoices (id, invoice_type, invoice_date, amount) VALUES "
            f"({row - 1}, {sql_text(invoice_types.get(raw_type, 'other'))}, {sql_text(iso(invoices.cell(row, 2).value))}, {n(invoices.cell(row, 3).value)});"
        )

    lines.append("")
    component_values = {
        "fixed_electricity_daily": prices["B2"].value,
        "fixed_water_daily": prices["B3"].value,
        "administration_daily": prices["B4"].value,
        "sunflowers_daily": prices["B5"].value,
    }
    fixed_rate = prices["D2"].value
    period_specs = [
        (1, "C", "D"),
        (2, "G", "H"),
    ]
    for period_id, actual_col, calculated_col in period_specs:
        fields = {
            "id": period_id,
            "name": summary[f"{actual_col}11"].value,
            "start_date": iso(summary[f"{actual_col}12"].value),
            "end_date": iso(summary[f"{actual_col}13"].value),
            "day_adjustment": -2,
            "actual_heating_rate": summary[f"{actual_col}14"].value,
            "actual_cooling_rate": summary[f"{actual_col}15"].value,
            "actual_water_rate": summary[f"{actual_col}16"].value,
            "actual_fixed_daily_rate": fixed_rate,
            "calculated_water_rate": summary[f"{calculated_col}16"].value,
            "calculated_fixed_daily_rate": fixed_rate,
            **component_values,
        }
        columns = ", ".join(fields.keys())
        values = ", ".join(sql_text(value) if isinstance(value, str) else n(value) for value in fields.values())
        lines.append(f"INSERT OR IGNORE INTO periods ({columns}) VALUES ({values});")

    lines.append("")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines), encoding="utf-8")
    print(f"Generated {output} with {len(dwelling_ids)} dwellings, {len(date_ids)} dates, and {invoices.max_row - 1} invoices.")


if __name__ == "__main__":
    main()
