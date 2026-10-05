import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def serialise(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    formula_book = load_workbook(source, data_only=False, read_only=False)
    value_book = load_workbook(source, data_only=True, read_only=False)

    report = {
        "source": str(source),
        "sheets": [],
        "defined_names": [
            {"name": item.name, "attr_text": item.attr_text}
            for item in formula_book.defined_names.values()
        ],
        "calculation": {
            "mode": formula_book.calculation.calcMode,
            "iterate": formula_book.calculation.iterate,
            "iterateCount": formula_book.calculation.iterateCount,
            "iterateDelta": formula_book.calculation.iterateDelta,
            "fullCalcOnLoad": formula_book.calculation.fullCalcOnLoad,
            "forceFullCalc": formula_book.calculation.forceFullCalc,
        },
    }

    for ws in formula_book.worksheets:
        cached_ws = value_book[ws.title]
        cells = []
        formulas = []
        data_types = Counter()
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cached = cached_ws[cell.coordinate].value
                record = {
                    "cell": cell.coordinate,
                    "value": serialise(cell.value),
                    "cached": serialise(cached),
                    "number_format": cell.number_format,
                    "style_id": cell.style_id,
                }
                cells.append(record)
                data_types[cell.data_type] += 1
                if cell.data_type == "f":
                    formulas.append(record)

        widths = {
            key: dim.width for key, dim in ws.column_dimensions.items() if dim.width
        }
        heights = {
            str(key): dim.height for key, dim in ws.row_dimensions.items() if dim.height
        }
        validations = []
        if ws.data_validations:
            for item in ws.data_validations.dataValidation:
                validations.append(
                    {
                        "sqref": str(item.sqref),
                        "type": item.type,
                        "formula1": item.formula1,
                        "formula2": item.formula2,
                    }
                )
        report["sheets"].append(
            {
                "title": ws.title,
                "state": ws.sheet_state,
                "dimensions": ws.calculate_dimension(),
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "merged_cells": [str(item) for item in ws.merged_cells.ranges],
                "formula_count": len(formulas),
                "data_types": dict(data_types),
                "tables": list(ws.tables.keys()),
                "auto_filter": str(ws.auto_filter.ref) if ws.auto_filter.ref else None,
                "data_validations": validations,
                "column_widths": widths,
                "row_heights": heights,
                "cells": cells,
                "formulas": formulas,
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "source": str(source),
        "output": str(output),
        "sheets": [
            {
                "title": item["title"],
                "dimensions": item["dimensions"],
                "formula_count": item["formula_count"],
                "cells": len(item["cells"]),
            }
            for item in report["sheets"]
        ],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
